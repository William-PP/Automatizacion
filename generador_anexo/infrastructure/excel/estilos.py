# -*- coding: utf-8 -*-
"""Capa de INFRAESTRUCTURA (excel): captura y aplicacion de estilos.

Las filas insertadas nacen sin estilo. Copiamos el estilo de una fila
"modelo" del anexo (col A blanca, B azul claro, C verde) a las celdas
donde escribimos valores. Se reconstruyen objetos nuevos (Font, Fill,
Border) para evitar copias de StyleProxy de openpyxl.
"""

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def capturar_estilo(ws, fila):
    est = []
    for col in range(1, 4):
        c = ws.cell(fila, col)
        f = c.font
        fill = c.fill
        b = c.border
        est.append(dict(
            font=Font(name=f.name, size=f.size, bold=f.bold, italic=f.italic,
                      underline=f.underline, color=f.color),
            fill=PatternFill(fill_type=fill.fill_type,
                             fgColor=fill.fgColor.rgb if fill.fgColor else None,
                             bgColor=fill.bgColor.rgb if fill.bgColor else None),
            border=Border(
                left=Side(style=getattr(getattr(b, "left", None), "style", None)),
                right=Side(style=getattr(getattr(b, "right", None), "style", None)),
                top=Side(style=getattr(getattr(b, "top", None), "style", None)),
                bottom=Side(style=getattr(getattr(b, "bottom", None), "style", None)),
            ),
            number_format=c.number_format,
            alignment=Alignment(horizontal=c.alignment.horizontal,
                                vertical=c.alignment.vertical,
                                wrap_text=c.alignment.wrap_text),
        ))
    return est


def aplicar_estilo(ws, fila, est):
    for col in range(1, 4):
        c = ws.cell(fila, col)
        s = est[col - 1]
        c.font = s["font"]
        c.fill = s["fill"]
        c.border = s["border"]
        c.number_format = s["number_format"]
        c.alignment = s["alignment"]


__all__ = ["capturar_estilo", "aplicar_estilo"]