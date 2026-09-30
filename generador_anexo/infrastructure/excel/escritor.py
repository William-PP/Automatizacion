# -*- coding: utf-8 -*-
"""Capa de INFRAESTRUCTURA (excel): escritura de valores en la hoja.

Funciones de bajo nivel que colocan valores/formulas y formato en la
hoja del anexo. Operan sobre una Regla (domain) y sus grupos ya
agrupados por domain.mapeo.
"""

from openpyxl.utils import column_index_from_string

from ..util.formato import fmt_valor, render_pct
from .estilos import aplicar_estilo


def escribir_total(ws, regla, total, fila_final, est):
    col = regla.col
    ic = column_index_from_string(col)
    celda = ws.cell(fila_final, ic)
    celda.value = fmt_valor(total)
    if regla.decimales:
        celda.number_format = "#,##0.00"
    if col == "B":
        f_c = regla.formula_c
        if f_c is not None:
            ws.cell(fila_final, 3).value = f_c.format(row=fila_final, factor=1)
    aplicar_estilo(ws, fila_final, est)


def escribir_filas(ws, regla, grupos, fila_inicio, est):
    ic = 2
    espejo = regla.espejo
    for i, (meta, valor) in enumerate(grupos):
        f = fila_inicio + i
        if regla.agregar == "detalle":
            ws.cell(f, ic).value = fmt_valor(valor)
            plantilla = regla.plantilla_detalle
            if plantilla:
                ws.cell(f, 1).value = plantilla.format(
                    matricula=meta.get("matricula", ""),
                    placa=meta.get("placa", ""),
                    valor=f"${int(round(valor)):,}",
                    pct=render_pct(meta.get("pct", 100)))
            else:
                ws.cell(f, 1).value = meta["tercero"]
            factor = (meta["pct"] / 100.0) if isinstance(meta, dict) else 1.0
            pct = meta["pct"] if isinstance(meta, dict) else 100.0
            ws.cell(f, 3).value = "=B{row}*{pct}%".format(
                row=f, pct=render_pct(pct))
        else:
            tercero = meta if isinstance(meta, str) else meta["tercero"]
            plantilla = regla.plantilla_tercero
            if plantilla:
                ws.cell(f, 1).value = plantilla.format(
                    tercero=tercero,
                    valor=f"${int(round(valor)):,}")
            else:
                ws.cell(f, 1).value = tercero
            celda = ws.cell(f, ic)
            celda.value = fmt_valor(valor)
            if espejo:
                ws.cell(f, 3).value = "=B{row}".format(row=f)
        aplicar_estilo(ws, f, est)
        if regla.decimales:
            ws.cell(f, 2).number_format = "#,##0.00"
            ws.cell(f, 3).number_format = "#,##0.00"


__all__ = ["escribir_total", "escribir_filas"]